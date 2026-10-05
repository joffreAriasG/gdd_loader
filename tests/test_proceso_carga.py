"""Orquestacion completa de una carga (procesar_archivo) sin SQL Server:
repositorio de control en memoria, staging simulado y merge simulado."""

from datetime import date
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook

from gdd_loader.domain import plantilla as pl
from gdd_loader.domain.sheet_config import SheetConfig
from gdd_loader.load.control_repository import CargaAnterior
from gdd_loader.pipeline import proceso_carga as pc
from gdd_loader.pipeline.merge_completo import ResumenMerge

HOJAS = [
    SheetConfig(
        nombre_hoja="DetalleAtributos",
        tabla_staging="staging.detalle_atributos",
        columnas=["codigo_dominio", "codigo_atributo"],
        columna_clave="codigo_dominio",
    ),
    SheetConfig(
        nombre_hoja="Respaldos",
        tabla_staging="staging.respaldos",
        columnas=["tipo_respaldo"],
        columna_clave="codigo_dominio",
        hereda_codigo_dominio_de="DetalleAtributos",
    ),
]
CONTRATO = {
    "DetalleAtributos": ["codigo_dominio", "codigo_atributo", "columna_no_leida"],
    "Respaldos": ["tipo_respaldo"],
}


def _version(estado="VIGENTE", version="1.0.0", columnas=CONTRATO):
    return pl.VersionPlantilla(
        "GDD-DOMINIO", version, estado, pl.calcular_hash_estructura(columnas), "0.2.0", columnas)


class RepoControlFalso:
    def __init__(self, versiones=(), ultima_ok=None):
        self._versiones = {v.version: v for v in versiones}
        self._ultima_ok = ultima_ok
        self.cargas: dict[int, dict] = {}
        self.detalles: list[tuple] = []

    def iniciar_carga(self, **campos):
        id_carga = len(self.cargas) + 1
        self.cargas[id_carga] = dict(campos)
        return id_carga

    def actualizar_carga(self, id_carga, **campos):
        self.cargas[id_carga].update(campos)

    def registrar_detalle(self, id_carga, etapa, objeto, **conteos):
        self.detalles.append((id_carga, etapa, objeto, conteos))

    def obtener_version(self, id_plantilla, version):
        return self._versiones.get(version)

    def versiones(self, id_plantilla):
        return list(self._versiones.values())

    def ultima_carga_ok(self, codigo_dominio):
        return self._ultima_ok


def _excel(path, dominio="TST", control=None, detalle_extra=True, autor="Autor Prueba"):
    wb = Workbook()
    ws = wb.active
    ws.title = "DetalleAtributos"
    enc = ["codigo_dominio", "codigo_atributo"] + (["columna_no_leida"] if detalle_extra else [])
    ws.append(enc)
    ws.append([dominio, "1"] + (["x"] if detalle_extra else []))
    ws.append([dominio, "2"] + (["y"] if detalle_extra else []))
    r = wb.create_sheet("Respaldos")
    r.append(["tipo_respaldo"])
    r.append(["Correo"])
    wb.create_sheet("Lista De Referencia").append(["id", "vals"])
    if control is not None:
        c = wb.create_sheet("_Control")
        c.append(["clave", "valor"])
        for k, v in control.items():
            c.append([k, v])
    wb.properties.lastModifiedBy = autor
    wb.save(path)
    return path


CONTROL_OK = {"id_plantilla": "GDD-DOMINIO", "version_plantilla": "1.0.0", "codigo_dominio": "TST"}


def _merge_ok(*_args):
    return [ResumenMerge("gdd.atributo", insertadas=2), ResumenMerge("gdd.respaldo", sin_cambios=1)]


@pytest.fixture
def entorno(tmp_path):
    entrada = tmp_path / "Entrada"
    entrada.mkdir()
    staging = MagicMock()
    staging.cargar.side_effect = lambda df, *a, **k: len(df)

    def construir(repo, exigir_control=False, funcion_merge=_merge_ok):
        return pc.ContextoCarga(
            control_repo=repo, staging_repo=staging, engine=object(), hojas=HOJAS,
            carpeta_archivo=tmp_path / "Archivo", carpeta_rechazados=tmp_path / "Rechazados",
            id_plantilla="GDD-DOMINIO", exigir_control=exigir_control, version_loader="0.2.0",
            equipo="EQUIPO", usuario_ejecucion="usuario", funcion_merge=funcion_merge,
            hoy=lambda: date(2026, 9, 25),
        )

    return entrada, staging, construir, tmp_path


def test_carga_completa_ok_con_control(entorno):
    entrada, staging, construir, tmp = entorno
    archivo = _excel(entrada / "cualquier nombre (1).xlsx", control=CONTROL_OK)
    repo = RepoControlFalso([_version()])

    r = pc.procesar_archivo(archivo, construir(repo))

    assert r.estado == pc.MERGE_OK and r.codigo_dominio == "TST" and r.id_carga == 1
    assert not archivo.exists()
    assert r.ruta_final.parent == tmp / "Archivo"
    assert r.ruta_final.name.startswith("TST_000001_")
    carga = repo.cargas[1]
    assert carga["estado"] == pc.MERGE_OK and carga["origen_version"] == "CONTROL"
    assert carga["version_plantilla"] == "1.0.0" and carga["codigo_dominio"] == "TST"
    assert carga["subido_por"] == "Autor Prueba" and carga["origen_subido_por"] == "DOCPROPS"
    assert carga["ruta_archivo_archivado"] == str(r.ruta_final)
    etapas = [(d[1], d[2]) for d in repo.detalles]
    assert etapas == [("STAGING", "DetalleAtributos"), ("STAGING", "Respaldos"),
                      ("MERGE", "gdd.atributo"), ("MERGE", "gdd.respaldo")]
    # cada DataFrame enviado a staging lleva el id_carga
    for llamada in staging.cargar.call_args_list:
        assert (llamada.args[0]["id_carga"] == 1).all()


def test_sin_control_y_exigido_se_rechaza_sin_tocar_staging(entorno):
    entrada, staging, construir, tmp = entorno
    archivo = _excel(entrada / "a.xlsx")
    repo = RepoControlFalso([_version()])

    r = pc.procesar_archivo(archivo, construir(repo, exigir_control=True))

    assert r.estado == pl.RECHAZADA_SIN_CONTROL
    assert r.ruta_final.parent == tmp / "Rechazados"
    assert r.ruta_final.name.startswith("SIN_DOMINIO_")
    staging.cargar.assert_not_called()


def test_transicion_sin_control_identifica_version_por_huella(entorno):
    entrada, _staging, construir, _tmp = entorno
    archivo = _excel(entrada / "a.xlsx")
    repo = RepoControlFalso([_version(estado="DEPRECADA", version="0.9.0",
                                      columnas={"DetalleAtributos": ["otra"]}), _version()])

    r = pc.procesar_archivo(archivo, construir(repo))

    assert r.estado == pc.MERGE_OK
    assert repo.cargas[1]["origen_version"] == "HUELLA"
    assert repo.cargas[1]["version_plantilla"] == "1.0.0"
    assert "Advertencias" in repo.cargas[1]["mensaje"]


def test_transicion_sin_control_y_sin_huella_conocida_se_rechaza(entorno):
    entrada, staging, construir, _tmp = entorno
    archivo = _excel(entrada / "a.xlsx", detalle_extra=False)
    repo = RepoControlFalso([_version()])

    r = pc.procesar_archivo(archivo, construir(repo))

    assert r.estado == pl.RECHAZADA_ESTRUCTURA
    assert "columna_no_leida" in r.mensaje
    staging.cargar.assert_not_called()


def test_version_retirada_se_rechaza(entorno):
    entrada, _staging, construir, _tmp = entorno
    archivo = _excel(entrada / "a.xlsx", control=CONTROL_OK)
    r = pc.procesar_archivo(archivo, construir(RepoControlFalso([_version(estado="RETIRADA")])))
    assert r.estado == pl.RECHAZADA_VERSION and "retirada" in r.mensaje


def test_control_de_otra_plantilla_se_rechaza(entorno):
    entrada, _staging, construir, _tmp = entorno
    archivo = _excel(entrada / "a.xlsx", control={**CONTROL_OK, "id_plantilla": "OTRA"})
    r = pc.procesar_archivo(archivo, construir(RepoControlFalso([_version()])))
    assert r.estado == pl.RECHAZADA_VERSION and "OTRA" in r.mensaje


def test_dominio_declarado_distinto_al_de_detalle_se_rechaza(entorno):
    entrada, staging, construir, _tmp = entorno
    archivo = _excel(entrada / "a.xlsx", control={**CONTROL_OK, "codigo_dominio": "ADS"})
    r = pc.procesar_archivo(archivo, construir(RepoControlFalso([_version()])))
    assert r.estado == pc.RECHAZADA_DOMINIO
    staging.cargar.assert_not_called()


def test_mismo_archivo_que_ultima_carga_ok_se_omite(entorno):
    entrada, staging, construir, tmp = entorno
    archivo = _excel(entrada / "a.xlsx", control=CONTROL_OK)
    from gdd_loader.extract.control_reader import hash_archivo
    repo = RepoControlFalso([_version()], ultima_ok=CargaAnterior(9, hash_archivo(archivo)))

    r = pc.procesar_archivo(archivo, construir(repo))

    assert r.estado == pc.OMITIDA_SIN_CAMBIOS and "9" in r.mensaje
    assert r.ruta_final.parent == tmp / "Archivo"
    staging.cargar.assert_not_called()


def test_error_de_staging_no_ejecuta_merge(entorno):
    entrada, staging, construir, tmp = entorno
    staging.cargar.side_effect = [2, RuntimeError("tabla no existe")]
    merge = MagicMock()
    archivo = _excel(entrada / "a.xlsx", control=CONTROL_OK)
    repo = RepoControlFalso([_version()])

    r = pc.procesar_archivo(archivo, construir(repo, funcion_merge=merge))

    assert r.estado == pc.ERROR_STAGING and "Respaldos" in r.mensaje
    merge.assert_not_called()
    assert r.ruta_final.parent == tmp / "Rechazados"
    assert repo.detalles[1][3]["error"] == "tabla no existe"


def test_merge_parcial_y_error_merge(entorno):
    entrada, _staging, construir, _tmp = entorno

    def parcial(*_a):
        return [ResumenMerge("gdd.atributo"), ResumenMerge("gdd.respaldo", error="catalogo")]

    def total(*_a):
        return [ResumenMerge("gdd.atributo", error="x"), ResumenMerge("gdd.respaldo", error="y")]

    a = _excel(entrada / "a.xlsx", control=CONTROL_OK)
    r1 = pc.procesar_archivo(a, construir(RepoControlFalso([_version()]), funcion_merge=parcial))
    b = _excel(entrada / "b.xlsx", control=CONTROL_OK)
    r2 = pc.procesar_archivo(b, construir(RepoControlFalso([_version()]), funcion_merge=total))

    assert r1.estado == pc.MERGE_PARCIAL and "gdd.respaldo: catalogo" in r1.mensaje
    assert r2.estado == pc.ERROR_MERGE


def test_excepcion_no_prevista_queda_registrada_como_error(entorno):
    entrada, _staging, construir, tmp = entorno

    def revienta(*_a):
        raise RuntimeError("conexion perdida")

    archivo = _excel(entrada / "a.xlsx", control=CONTROL_OK)
    repo = RepoControlFalso([_version()])
    r = pc.procesar_archivo(archivo, construir(repo, funcion_merge=revienta))

    assert r.estado == pc.ERROR and "conexion perdida" in r.mensaje
    assert repo.cargas[1]["estado"] == pc.ERROR
    assert r.ruta_final.parent == tmp / "Rechazados"


def test_si_no_se_puede_mover_se_conserva_el_estado(entorno, monkeypatch):
    entrada, _staging, construir, _tmp = entorno
    archivo = _excel(entrada / "a.xlsx", control=CONTROL_OK)
    repo = RepoControlFalso([_version()])

    def falla(*_a):
        raise PermissionError("bloqueado por OneDrive")

    monkeypatch.setattr(pc, "mover", falla)
    r = pc.procesar_archivo(archivo, construir(repo))

    assert r.estado == pc.MERGE_OK and archivo.exists()
    assert "No se pudo archivar" in repo.cargas[1]["mensaje"]
    assert "ruta_archivo_archivado" not in repo.cargas[1]


# --- Notificacion (v0.5.0) -------------------------------------------------------


class NotificadorFalso:
    def __init__(self, fotos=(), falla_instantanea=False, falla_notificar=False):
        self.fotos = list(fotos)
        self.falla_instantanea = falla_instantanea
        self.falla_notificar = falla_notificar
        self.instantaneas: list[str] = []
        self.notificados = []

    def instantanea(self, dominio):
        self.instantaneas.append(dominio)
        if self.falla_instantanea:
            raise RuntimeError("sin conexion")
        return self.fotos.pop(0)

    def notificar(self, resultado):
        if self.falla_notificar:
            raise OSError("carpeta no disponible")
        self.notificados.append(resultado)
        return "ruta/evento.json"


def _fotos_con_un_atributo_nuevo():
    from gdd_loader.notificacion.instantanea import Registro
    return [{"Atributos": {}}, {"Atributos": {("1",): Registro("1 – Poliza")}}]


def test_notificacion_toma_foto_antes_y_despues_del_merge(entorno):
    entrada, _staging, construir, _tmp = entorno
    archivo = _excel(entrada / "a.xlsx", control=CONTROL_OK)
    orden = []
    notif = NotificadorFalso(_fotos_con_un_atributo_nuevo())
    instantanea_original = notif.instantanea
    notif.instantanea = lambda d: (orden.append("foto"), instantanea_original(d))[1]

    def merge(*args):
        orden.append("merge")
        return _merge_ok()

    ctx = construir(RepoControlFalso([_version()]), funcion_merge=merge)
    ctx.notificador = notif
    r = pc.procesar_archivo(archivo, ctx)

    assert r.estado == pc.MERGE_OK
    assert orden == ["foto", "merge", "foto"]
    assert [c.etiqueta for c in r.cambios[0].nuevos] == ["1 – Poliza"]
    assert notif.notificados == [r]
    assert r.version_plantilla == "1.0.0" and r.subido_por == "Autor Prueba"
    assert r.ruta_notificacion == "ruta/evento.json"
    assert r.ruta_final is not None  # se notifica despues de archivar


def test_notificacion_de_rechazo_sin_fotos(entorno):
    entrada, _staging, construir, _tmp = entorno
    archivo = _excel(entrada / "a.xlsx")
    notif = NotificadorFalso()
    ctx = construir(RepoControlFalso([_version()]), exigir_control=True)
    ctx.notificador = notif

    r = pc.procesar_archivo(archivo, ctx)

    assert r.estado == pl.RECHAZADA_SIN_CONTROL
    assert notif.instantaneas == [] and r.cambios is None
    assert notif.notificados == [r]


def test_falla_de_foto_no_afecta_la_carga(entorno):
    entrada, _staging, construir, _tmp = entorno
    archivo = _excel(entrada / "a.xlsx", control=CONTROL_OK)
    notif = NotificadorFalso(falla_instantanea=True)
    ctx = construir(RepoControlFalso([_version()]))
    ctx.notificador = notif

    r = pc.procesar_archivo(archivo, ctx)

    assert r.estado == pc.MERGE_OK
    assert notif.instantaneas == ["TST"]  # no reintenta la foto de despues
    assert r.cambios is None and "No se pudo calcular el detalle" in r.motivo_sin_detalle
    assert notif.notificados == [r]


def test_falla_al_notificar_no_afecta_la_carga(entorno):
    entrada, _staging, construir, _tmp = entorno
    archivo = _excel(entrada / "a.xlsx", control=CONTROL_OK)
    repo = RepoControlFalso([_version()])
    ctx = construir(repo)
    ctx.notificador = NotificadorFalso(_fotos_con_un_atributo_nuevo(), falla_notificar=True)

    r = pc.procesar_archivo(archivo, ctx)

    assert r.estado == pc.MERGE_OK and repo.cargas[1]["estado"] == pc.MERGE_OK
    assert r.ruta_notificacion is None
