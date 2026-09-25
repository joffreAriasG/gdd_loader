"""Piezas de la Fase 1 que tocan codigo existente: configuracion, bitacora
con id_carga, id_carga en staging, resumen del merge completo y la
herramienta de registro de plantillas."""

from unittest.mock import MagicMock

import pandas as pd
import pytest
from openpyxl import Workbook

from gdd_loader.config import Settings
from gdd_loader.domain import plantilla as pl
from gdd_loader.domain.sheet_config import HOJAS_ADS, SheetConfig
from gdd_loader.load import gdd_repository as repo
from gdd_loader.pipeline import merge_completo
from gdd_loader.pipeline.carga_staging import ejecutar
from gdd_loader.plantilla_cli import construir_version


def _settings(tmp_path, **cambios):
    base = dict(
        carpeta_origen=tmp_path / "Entrada", sql_servidor="srv", sql_base_datos="DB",
        sql_driver="d", sql_auth_modo="WINDOWS", sql_usuario="", sql_password="",
        estrategia_staging="DOMINIO", log_dir=tmp_path / "logs", log_level="INFO",
        carpeta_archivo=tmp_path / "Archivo", carpeta_rechazados=tmp_path / "Rechazados",
    )
    base.update(cambios)
    return Settings(**base)


def test_validar_procesamiento_ok(tmp_path):
    _settings(tmp_path).validar_procesamiento()


@pytest.mark.parametrize("campo, texto", [
    ("sin_archivo", "GDD_CARPETA_ARCHIVO"),
    ("truncate", "DOMINIO"),
    ("misma_carpeta", "distintas"),
])
def test_validar_procesamiento_errores(tmp_path, campo, texto):
    cambios = {
        "sin_archivo": {"carpeta_archivo": None},
        "truncate": {"estrategia_staging": "TRUNCATE"},
        "misma_carpeta": {"carpeta_rechazados": tmp_path / "Archivo"},
    }[campo]
    with pytest.raises(ValueError, match=texto):
        _settings(tmp_path, **cambios).validar_procesamiento()


def test_registrar_bitacora_con_y_sin_id_carga():
    conn = MagicMock()
    repo.registrar_bitacora(conn, "atributo", 1, "INSERTAR", "ADS", "ADS-1", "1")
    sql_sin = str(conn.execute.call_args.args[0])
    assert "id_carga" not in sql_sin and "id_carga" not in conn.execute.call_args.args[1]

    repo.registrar_bitacora(conn, "atributo", 1, "INSERTAR", "ADS", "ADS-1", "1", id_carga=77)
    sql_con = str(conn.execute.call_args.args[0])
    assert "id_carga" in sql_con and conn.execute.call_args.args[1]["id_carga"] == 77


def test_staging_sin_id_carga_no_agrega_columna(tmp_path):
    archivo = tmp_path / "a.xlsx"
    with pd.ExcelWriter(archivo, engine="openpyxl") as w:
        pd.DataFrame({"codigo_dominio": ["A"]}).to_excel(w, sheet_name="DetalleAtributos", index=False)
    cfg = SheetConfig("DetalleAtributos", "staging.t", ["codigo_dominio"], columna_clave="codigo_dominio")
    staging = MagicMock()
    staging.cargar.return_value = 1

    ejecutar(archivo, [cfg], staging, "DOMINIO")
    assert "id_carga" not in staging.cargar.call_args.args[0].columns

    ejecutar(archivo, [cfg], staging, "DOMINIO", id_carga=5)
    assert list(staging.cargar.call_args.args[0]["id_carga"]) == [5]


def test_mergear_dominio_completo_normaliza_resultados(monkeypatch):
    from gdd_loader.pipeline.carga_gdd import ResultadoMergeDominio, ResultadoMergeInvestigacion

    llamadas = {}

    def dominio(engine, staging, codigo, id_carga=None):
        llamadas["id_carga"] = id_carga
        return ResultadoMergeDominio(codigo, atributos_insertados=3, atributos_omitidos_fecha_retrocede=1,
                                     fuentes_reemplazadas=2, consumos_eliminados=4)

    def simple_ok(engine, staging, codigo):
        return ResultadoMergeInvestigacion(codigo, insertadas=1, sin_cambios=5)

    def simple_error(engine, staging, codigo):
        return ResultadoMergeInvestigacion(codigo, error="catalogo no encontrado")

    monkeypatch.setattr(merge_completo, "ejecutar_merge_dominio", dominio)
    monkeypatch.setattr(merge_completo, "ejecutar_merge_investigacion_dominio", simple_ok)
    monkeypatch.setattr(merge_completo, "ejecutar_merge_estructura_dominio", simple_ok)
    monkeypatch.setattr(merge_completo, "ejecutar_merge_respaldos_dominio", simple_ok)
    monkeypatch.setattr(merge_completo, "ejecutar_merge_plan_remediacion_dominio", simple_error)

    res = {r.objeto: r for r in merge_completo.mergear_dominio_completo(None, None, "ADS", id_carga=9)}

    assert llamadas["id_carga"] == 9
    assert list(res) == ["gdd.atributo", "gdd.atributo_fuente_oficial", "gdd.atributo_fuente_consumo",
                         "gdd.investigacion", "gdd.dominio_responsable", "gdd.respaldo",
                         "gdd.plan_remediacion"]
    assert res["gdd.atributo"].insertadas == 3 and res["gdd.atributo"].omitidas == 1
    assert res["gdd.atributo_fuente_oficial"].reemplazadas == 2
    assert res["gdd.atributo_fuente_consumo"].eliminadas == 4
    assert res["gdd.investigacion"].sin_cambios == 5
    assert not res["gdd.plan_remediacion"].ok


def _molde(path, version="1.0.0", quitar=None):
    wb = Workbook()
    wb.remove(wb.active)
    for cfg in HOJAS_ADS:
        ws = wb.create_sheet(cfg.nombre_hoja)
        cols = [c for c in cfg.columnas if c != quitar] + ["columna_extra_no_leida"]
        # las hojas que heredan codigo_dominio no lo traen en el Excel
        ws.append(cols)
    wb.create_sheet("Lista De Referencia").append(["id", "vals"])
    c = wb.create_sheet("_Control")
    c.append(["clave", "valor"])
    c.append(["id_plantilla", "GDD-DOMINIO"])
    c.append(["version_plantilla", version])
    wb.save(path)
    return path


def test_construir_version_desde_molde(tmp_path):
    hojas = [c.nombre_hoja for c in HOJAS_ADS]
    v, problemas = construir_version(_molde(tmp_path / "m.xlsx"), "GDD-DOMINIO", "1.0.0", "0.2.0", hojas)
    assert problemas == []
    assert set(v.columnas) == set(hojas)
    assert "Lista De Referencia" not in v.columnas
    assert v.hash_estructura == pl.calcular_hash_estructura(v.columnas)
    assert v.columnas["Respaldos"][-1] == "columna_extra_no_leida"


def test_construir_version_detecta_problemas(tmp_path):
    hojas = [c.nombre_hoja for c in HOJAS_ADS]
    molde = _molde(tmp_path / "m.xlsx", version="0.9.0", quitar="enlace_respaldo")
    _v, problemas = construir_version(molde, "GDD-DOMINIO", "1.0.0", "0.2.0", hojas)
    assert any("enlace_respaldo" in p for p in problemas)
    assert any("_Control del molde declara" in p for p in problemas)
