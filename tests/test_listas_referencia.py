"""Listas de referencia (Fase 2): lectura del molde, comparacion con la BD,
modo BD del generador, semilla para pruebas y limpieza de Reporte_Errores.
Datos sinteticos."""

from pathlib import Path
from unittest.mock import MagicMock

from openpyxl import Workbook, load_workbook

from gdd_loader.exportar import listas_referencia as lr
from gdd_loader.exportar.escritor_molde import llenar_molde


def _molde(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "DetalleAtributos"
    ws.append(["codigo_dominio", "atributo"])
    listas = wb.create_sheet("ListaDeReferencia")
    listas.append(["id", "vals"])
    for fila in [("Criticidad", "Crítico"), ("Criticidad", "No crítico"), ("Criticidad", "No crítico"),
                 ("DatoPersonal", "Sí"), ("DatoPersonal", "No"),
                 ("Sensibilidad", "1. Bajo"), ("Sensibilidad", "2. Medio"), (None, None)]:
        listas.append(list(fila))
    consumo = wb.create_sheet("FuentesConsumo")
    consumo.append(["coleccion", "name"])
    consumo.append(["col_a", "tabla_1"])
    consumo.append(["col_a", "tabla_2"])
    errores = wb.create_sheet("Reporte_Errores")
    errores.append(["Atributo", "Sección", "FilaExcel", "Campo", "Detalle"])
    errores.append(["Atributo real", "DetalleAtributos", 14, "sensibilidad", "detalle del error"])
    wb.save(path)
    return path


class RepoListas:
    def __init__(self, simples=None, dato_personal=None, consumo=None):
        self.simples = simples or {}
        self.dp = dato_personal or {}
        self.consumo = consumo if consumo is not None else []

    def valores_simples(self, tabla, col):
        return self.simples.get(tabla, [])

    def valores_dato_personal(self, cod, desc):
        return self.dp.get(cod, [])

    def fuentes_consumo(self):
        return self.consumo


def test_leer_listas_molde_agrupa_sin_duplicados(tmp_path):
    listas = lr.leer_listas_molde(_molde(tmp_path / "m.xlsx"))
    assert listas.grupos == {"Criticidad": ["Crítico", "No crítico"], "DatoPersonal": ["Sí", "No"],
                             "Sensibilidad": ["1. Bajo", "2. Medio"]}
    assert listas.consumo == [("col_a", "tabla_1"), ("col_a", "tabla_2")]


def test_comparar_informa_faltantes_y_sobrantes(tmp_path):
    listas = lr.leer_listas_molde(_molde(tmp_path / "m.xlsx"))
    repo = RepoListas(simples={"cat_criticidad": ["Crítico", "Medio"]},
                      dato_personal={"codigo_sensibilidad": ["1. Bajo", "2. Medio"]},
                      consumo=[("col_a", "tabla_1")])
    difs = {d.grupo: d for d in lr.comparar(listas, repo)}
    assert difs["Criticidad"].faltan_en_bd == ["No crítico"] and difs["Criticidad"].solo_en_bd == ["Medio"]
    assert difs["Sensibilidad"].coincide
    assert difs["DatoPersonal"].coincide and "sin catalogo" in difs["DatoPersonal"].origen_bd
    assert difs["FuentesConsumo"].faltan_en_bd == ["col_a / tabla_2"]


def test_construir_listas_modo_bd_con_respaldo_del_molde(tmp_path):
    listas = lr.leer_listas_molde(_molde(tmp_path / "m.xlsx"))
    repo = RepoListas(simples={"cat_criticidad": ["Alto", "Bajo"]}, consumo=[])
    filas, avisos = lr.construir_listas(listas, repo)
    assert filas["ListaDeReferencia"] == [
        ["Criticidad", "Alto"], ["Criticidad", "Bajo"],  # desde la BD
        ["DatoPersonal", "Sí"], ["DatoPersonal", "No"],  # sin catalogo: molde
        ["Sensibilidad", "1. Bajo"], ["Sensibilidad", "2. Medio"],  # catalogo vacio: molde
    ]
    assert filas["FuentesConsumo"] == [["col_a", "tabla_1"], ["col_a", "tabla_2"]]
    assert any("Sensibilidad" in a for a in avisos) and any("FuentesConsumo" in a for a in avisos)


def test_script_semilla_solo_catalogos_simples_y_escapa_comillas(tmp_path):
    listas = lr.ListasMolde({"Criticidad": ["No crítico", "O'Brien"], "Sensibilidad": ["3. Alto"]}, [])
    difs = [lr.Diferencia("Criticidad", "x", ["No crítico", "O'Brien"], []),
            lr.Diferencia("Sensibilidad", "y", ["3. Alto"], [])]
    sql = lr.script_semilla(listas, difs)
    assert "INSERT INTO gdd.cat_criticidad (nombre_criticidad) VALUES (N'No crítico');" in sql
    assert "N'O''Brien'" in sql
    assert "IsIdentity" in sql and "BEGIN TRANSACTION" in sql and "COMMIT" in sql
    assert "cat_dato_personal" not in sql.split("/* No incluidos")[0]
    assert "Sensibilidad: 1 valor(es)" in sql


def test_limpiar_reporte_errores_sin_control(tmp_path):
    molde = _molde(tmp_path / "m.xlsx")
    salida = tmp_path / "limpio.xlsx"
    llenar_molde(molde, salida, {"Reporte_Errores": []}, None)
    wb = load_workbook(salida)
    assert "_Control" not in wb.sheetnames
    ws = wb["Reporte_Errores"]
    assert [c.value for c in ws[1]] == ["Atributo", "Sección", "FilaExcel", "Campo", "Detalle"]
    assert all(c.value is None for c in ws[2])
    assert wb["ListaDeReferencia"]["B2"].value == "Crítico"  # el resto intacto


def test_listas_repository_consulta_solo_lectura():
    engine = MagicMock()
    conn = engine.connect.return_value.__enter__.return_value
    conn.execute.return_value.all.return_value = [("6", "Crediticios"), ("2", "Identificativos"), ("2", "Identificativos")]
    repo = lr.ListasRepository(engine)
    assert repo.valores_dato_personal("codigo_categoria_nivel_uno", "categoria_nivel_uno") == \
        ["2. Identificativos", "6. Crediticios"]
    sql = str(conn.execute.call_args.args[0])
    assert sql.startswith("SELECT DISTINCT") and "gdd.cat_dato_personal" in sql
