"""Fase 2: escritor de molde, fuentes de datos, generador y validacion de
plantilla desactualizada. Moldes sinteticos (sin datos reales)."""

import shutil
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from gdd_loader.domain import plantilla as pl
from gdd_loader.domain.sheet_config import SheetConfig
from gdd_loader.exportar import fuente_datos as fd
from gdd_loader.exportar.escritor_molde import MoldeInvalidoError, llenar_molde
from gdd_loader.exportar.generador import (
    ContextoGenerador,
    GeneracionError,
    buscar_molde,
    generar_plantilla,
)
from gdd_loader.extract.control_reader import leer_control, leer_estructura
from gdd_loader.load.control_repository import CargaAnterior

ENCABEZADOS = {
    "DetalleAtributos": ["codigo_dominio", "codigo_atributo", "atributo", "fecha_aprobada"],
    "Respaldos": ["tipo_respaldo", "enlace_respaldo"],
}
PARTE_SCRIPTS = "customXml/item1.xml"
CONTENIDO_SCRIPTS = b'<scriptIds xmlns="urn:prueba"><scriptId id="script-validacion"/></scriptIds>'


def _crear_molde(path: Path, filas: dict | None = None, control: dict | None = None) -> Path:
    """Molde con tablas de Excel, una hoja auxiliar, formato de fecha en
    DetalleAtributos.fecha_aprobada y una parte customXml (como la referencia
    a los Office Scripts del molde real)."""
    wb = Workbook()
    wb.remove(wb.active)
    for i, (hoja, cols) in enumerate(ENCABEZADOS.items(), start=1):
        ws = wb.create_sheet(hoja)
        ws.append(cols)
        datos = (filas or {}).get(hoja, [])
        for f in datos:
            ws.append(f)
        # filas en blanco con estilo, como el molde real
        ultima = max(len(datos) + 1, 5)
        for r in range(len(datos) + 2, ultima + 1):
            for c in range(1, len(cols) + 1):
                ws.cell(r, c).number_format = "yyyy-mm-dd" if cols[c - 1] == "fecha_aprobada" else "@"
        letra = chr(64 + len(cols))
        ws.add_table(Table(displayName=f"tbl_{i}", ref=f"A1:{letra}{ultima}"))
    wb.create_sheet("ListaDeReferencia").append(["id", "vals"])
    if control:
        c = wb.create_sheet("_Control")
        c.append(["clave", "valor"])
        for k, v in control.items():
            c.append([k, v])
    wb.save(path)
    # agregar customXml como lo trae el molde real
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w") as zout:
        for info in zin.infolist():
            zout.writestr(info, zin.read(info.filename))
        zout.writestr(PARTE_SCRIPTS, CONTENIDO_SCRIPTS)
    tmp.replace(path)
    return path


def _version(molde: Path, estado="VIGENTE", version="1.0.0") -> pl.VersionPlantilla:
    cols = pl.restringir_a_contrato(leer_estructura(molde), ENCABEZADOS)
    return pl.VersionPlantilla("GDD-DOMINIO", version, estado, pl.calcular_hash_estructura(cols), "0.3.0", cols)


# --- escritor ----------------------------------------------------------------


def test_llenar_molde_preserva_partes_y_escribe_datos(tmp_path):
    molde = _crear_molde(tmp_path / "molde.xlsx")
    destino = tmp_path / "salida" / "gen.xlsx"
    datos = {
        "DetalleAtributos": [
            ["TST", 1, "Nombre <con> & especiales", date(2026, 1, 31)],
            ["TST", 2, "  espacios  ", datetime(2026, 2, 1, 0, 0)],
            ["TST", Decimal("3"), None, None],
        ],
        "Respaldos": [],
    }
    r = llenar_molde(molde, destino, datos, {"id_plantilla": "GDD-DOMINIO", "version_plantilla": "1.0.0"})

    assert r.filas_por_hoja == {"DetalleAtributos": 3, "Respaldos": 0}
    assert r.advertencias == []
    with zipfile.ZipFile(destino) as z:
        assert z.read(PARTE_SCRIPTS) == CONTENIDO_SCRIPTS
    wb = load_workbook(destino)
    ws = wb["DetalleAtributos"]
    assert [c.value for c in ws[2]] == ["TST", 1, "Nombre <con> & especiales", datetime(2026, 1, 31)]
    assert ws["C3"].value == "  espacios  "
    assert ws["B4"].value == 3 and ws["C4"].value is None
    assert ws.tables["tbl_1"].ref == "A1:D4"
    assert wb["Respaldos"].tables["tbl_2"].ref == "A1:B2"  # tabla vacia conserva una fila
    assert wb["_Control"].sheet_state == "veryHidden"
    assert leer_control(destino) == {"id_plantilla": "GDD-DOMINIO", "version_plantilla": "1.0.0"}
    assert wb.sheetnames[-1] == "_Control" and "ListaDeReferencia" in wb.sheetnames
    # misma huella que el molde
    assert pl.calcular_hash_estructura(pl.restringir_a_contrato(leer_estructura(destino), ENCABEZADOS)) == \
        _version(molde).hash_estructura


def test_fecha_en_columna_sin_formato_de_fecha_se_escribe_como_texto_legible_por_el_loader(tmp_path):
    molde = _crear_molde(tmp_path / "molde.xlsx")
    destino = tmp_path / "gen.xlsx"
    r = llenar_molde(molde, destino, {"Respaldos": [["Correo", date(2026, 9, 15)]]}, {})
    assert "no tiene formato de fecha" in r.advertencias[0]
    assert load_workbook(destino)["Respaldos"]["B2"].value == "2026-09-15 00:00:00"


def test_molde_con_control_existente_se_actualiza_sin_duplicar(tmp_path):
    molde = _crear_molde(tmp_path / "molde.xlsx", control={"id_plantilla": "GDD-DOMINIO", "version_plantilla": "1.0.1"})
    destino = tmp_path / "gen.xlsx"
    llenar_molde(molde, destino, {}, {"id_plantilla": "GDD-DOMINIO", "version_plantilla": "1.0.1", "codigo_dominio": "TST"})
    wb = load_workbook(destino)
    assert wb.sheetnames.count("_Control") == 1
    assert leer_control(destino)["codigo_dominio"] == "TST"


def test_fila_con_mas_columnas_que_el_molde_falla(tmp_path):
    molde = _crear_molde(tmp_path / "molde.xlsx")
    with pytest.raises(MoldeInvalidoError, match="columnas"):
        llenar_molde(molde, tmp_path / "g.xlsx", {"Respaldos": [["a", "b", "c"]]}, {})


# --- fuentes de datos -----------------------------------------------------------


def test_alinear_a_molde_por_nombre_de_columna():
    datos = {"Respaldos": [{"enlace_respaldo": "http://x", "tipo_respaldo": "Correo", "vieja": 1}]}
    filas, avisos = fd.alinear_a_molde(datos, {"Respaldos": ["tipo_respaldo", "enlace_respaldo", "nueva"]})
    assert filas == {"Respaldos": [["Correo", "http://x", None]]}
    assert any("nueva" in a for a in avisos) and any("vieja" in a for a in avisos)


def test_datos_desde_staging_ordena_y_exige_misma_carga():
    cfg = SheetConfig("MetadataTecnica", "staging.metadata_tecnica", ["codigo_dominio_atributo", "clase"],
                      columna_clave="codigo_dominio_atributo", columna_clave_prefijo_dominio=True)
    repo = MagicMock()
    repo.leer_por_prefijo.return_value = [
        {"codigo_dominio_atributo": "TST-10", "clase": "b", "id_carga": 7},
        {"codigo_dominio_atributo": "TST-9", "clase": "a", "id_carga": 7},
    ]
    datos = fd.datos_desde_staging(repo, [cfg], "TST", 7)
    assert [f["codigo_dominio_atributo"] for f in datos["MetadataTecnica"]] == ["TST-9", "TST-10"]
    assert "id_carga" not in datos["MetadataTecnica"][0]
    with pytest.raises(fd.FuenteNoDisponibleError, match="otra carga"):
        fd.datos_desde_staging(repo, [cfg], "TST", 8)


# --- generador --------------------------------------------------------------------


class RepoFalso:
    def __init__(self, versiones, ultima):
        self._v, self._u = versiones, ultima

    def versiones(self, _p):
        return self._v

    def ultima_carga_ok(self, _d):
        return self._u


@pytest.fixture
def escenario(tmp_path):
    moldes = tmp_path / "Moldes"
    moldes.mkdir()
    molde = _crear_molde(moldes / "PlantillaV1.xlsx")
    ok = _crear_molde(tmp_path / "Archivo" / "TST_000005_aa.xlsx" if (tmp_path / "Archivo").mkdir() is None else None,
                      filas={"DetalleAtributos": [["TST", 1, "Uno", date(2026, 1, 1)],
                                                  ["TST", 2, "Dos", date(2026, 1, 2)]],
                             "Respaldos": [["Correo", "http://evidencia"]]})
    # una plantilla ya llenada en la carpeta de moldes NO debe confundirse con el molde
    shutil.copy(ok, moldes / "TST_llenado.xlsx")
    version = _version(molde)
    return tmp_path, molde, ok, version


def test_generar_desde_archivo_de_la_ultima_carga_ok(escenario):
    tmp, molde, ok, version = escenario
    ctx = ContextoGenerador(RepoFalso([version], CargaAnterior(5, "h", str(ok))), MagicMock(), [],
                            "GDD-DOMINIO", tmp / "Moldes", tmp / "Plantillas")

    r = generar_plantilla("TST", ctx, datetime(2026, 9, 25, 12, 0))

    assert r.ruta.name == "TST_v1.0.0_c000005.xlsx" and r.fuente == "ARCHIVO"
    assert r.filas_por_hoja == {"DetalleAtributos": 2, "Respaldos": 1}
    control = leer_control(r.ruta)
    assert control["codigo_dominio"] == "TST" and control["id_carga_base"] == "5"
    assert control["version_plantilla"] == "1.0.0" and control["id_envio"] == r.id_envio
    ws = load_workbook(r.ruta)["DetalleAtributos"]
    assert [ws.cell(3, c).value for c in (1, 2, 3)] == ["TST", 2, "Dos"]
    with zipfile.ZipFile(r.ruta) as z:
        assert z.read(PARTE_SCRIPTS) == CONTENIDO_SCRIPTS


def test_molde_con_filas_predefinidas_fuera_de_la_hoja_ancla_sigue_siendo_molde(tmp_path):
    molde = _crear_molde(tmp_path / "m.xlsx", filas={"Respaldos": [["Correo", None]]})
    assert buscar_molde(tmp_path, _version(molde)) == molde


def test_buscar_molde_ignora_plantillas_llenadas_y_detecta_ambiguedad(escenario):
    tmp, molde, _ok, version = escenario
    assert buscar_molde(tmp / "Moldes", version) == molde
    shutil.copy(molde, tmp / "Moldes" / "copia_en_blanco.xlsx")
    with pytest.raises(GeneracionError, match="2 moldes en blanco"):
        buscar_molde(tmp / "Moldes", version)


def test_generar_sin_carga_ok_o_sin_vigente_falla(escenario):
    tmp, _molde, _ok, version = escenario
    base = dict(staging_repo=MagicMock(), hojas=[], id_plantilla="GDD-DOMINIO",
                carpeta_moldes=tmp / "Moldes", carpeta_salida=tmp / "P")
    with pytest.raises(GeneracionError, match="ninguna carga MERGE_OK"):
        generar_plantilla("TST", ContextoGenerador(control_repo=RepoFalso([version], None), **base))
    with pytest.raises(GeneracionError, match="VIGENTE"):
        generar_plantilla("TST", ContextoGenerador(control_repo=RepoFalso([], None), **base))


def test_generar_desde_staging_si_el_archivo_no_existe(escenario):
    tmp, _molde, _ok, version = escenario
    staging = MagicMock()
    staging.leer.side_effect = lambda tabla, _col, _dom: {
        "staging.detalle_atributos": [{"codigo_dominio": "TST", "codigo_atributo": 1, "atributo": "Uno",
                                       "fecha_aprobada": date(2026, 1, 1), "id_carga": 5}],
        "staging.respaldos": [],
    }[tabla]
    hojas = [SheetConfig("DetalleAtributos", "staging.detalle_atributos", ENCABEZADOS["DetalleAtributos"],
                         columna_clave="codigo_dominio"),
             SheetConfig("Respaldos", "staging.respaldos", ENCABEZADOS["Respaldos"], columna_clave="codigo_dominio")]
    ctx = ContextoGenerador(RepoFalso([version], CargaAnterior(5, "h", str(tmp / "no_existe.xlsx"))), staging,
                            hojas, "GDD-DOMINIO", tmp / "Moldes", tmp / "Plantillas")

    r = generar_plantilla("TST", ctx)

    assert r.fuente == "STAGING" and r.filas_por_hoja["DetalleAtributos"] == 1
    assert any("staging" in a for a in r.advertencias)


# --- validacion de plantilla desactualizada ------------------------------------------


def _ctx_proceso(tmp_path, anterior):
    from gdd_loader.pipeline import proceso_carga as pc
    from gdd_loader.pipeline.merge_completo import ResumenMerge

    class Repo:
        cargas = {}

        def iniciar_carga(self, **c):
            self.cargas[1] = dict(c)
            return 1

        def actualizar_carga(self, i, **c):
            self.cargas[i].update(c)

        def registrar_detalle(self, *a, **k):
            pass

        def obtener_version(self, _p, _v):
            return VERSION_PROCESO

        def ultima_carga_ok(self, _d):
            return anterior

    staging = MagicMock()
    staging.cargar.side_effect = lambda df, *a, **k: len(df)
    hojas = [SheetConfig("DetalleAtributos", "staging.detalle_atributos", ["codigo_dominio", "codigo_atributo"],
                         columna_clave="codigo_dominio")]
    return pc, pc.ContextoCarga(
        control_repo=Repo(), staging_repo=staging, engine=None, hojas=hojas,
        carpeta_archivo=tmp_path / "A", carpeta_rechazados=tmp_path / "R", id_plantilla="GDD-DOMINIO",
        exigir_control=False, version_loader="0.3.0", equipo="E", usuario_ejecucion="u",
        funcion_merge=lambda *a: [ResumenMerge("gdd.atributo")], hoy=lambda: date(2026, 9, 25))


VERSION_PROCESO = pl.VersionPlantilla(
    "GDD-DOMINIO", "1.0.0", "VIGENTE",
    pl.calcular_hash_estructura({"DetalleAtributos": ["codigo_dominio", "codigo_atributo"]}),
    "0.3.0", {"DetalleAtributos": ["codigo_dominio", "codigo_atributo"]})


def _archivo_con_base(tmp_path, id_carga_base):
    wb = Workbook()
    ws = wb.active
    ws.title = "DetalleAtributos"
    ws.append(["codigo_dominio", "codigo_atributo"])
    ws.append(["TST", "1"])
    c = wb.create_sheet("_Control")
    c.append(["clave", "valor"])
    for k, v in {"id_plantilla": "GDD-DOMINIO", "version_plantilla": "1.0.0", "codigo_dominio": "TST",
                 "id_carga_base": id_carga_base}.items():
        if v is not None:
            c.append([k, v])
    p = tmp_path / f"p_{id_carga_base}.xlsx"
    wb.save(p)
    return p


def test_plantilla_generada_sobre_la_ultima_carga_se_acepta(tmp_path):
    pc, ctx = _ctx_proceso(tmp_path, CargaAnterior(5, "otro-hash"))
    r = pc.procesar_archivo(_archivo_con_base(tmp_path, 5), ctx)
    assert r.estado == pc.MERGE_OK


def test_plantilla_generada_sobre_carga_anterior_se_rechaza(tmp_path):
    pc, ctx = _ctx_proceso(tmp_path, CargaAnterior(6, "otro-hash"))
    r = pc.procesar_archivo(_archivo_con_base(tmp_path, 5), ctx)
    assert r.estado == pc.RECHAZADA_DESACTUALIZADA
    assert "carga 5" in r.mensaje and "es la 6" in r.mensaje
    ctx.staging_repo.cargar.assert_not_called()


def test_plantilla_sin_id_carga_base_se_acepta_con_advertencia(tmp_path):
    pc, ctx = _ctx_proceso(tmp_path, CargaAnterior(6, "otro-hash"))
    r = pc.procesar_archivo(_archivo_con_base(tmp_path, None), ctx)
    assert r.estado == pc.MERGE_OK
    assert any("id_carga_base" in a for a in r.advertencias)
