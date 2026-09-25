import hashlib

from openpyxl import Workbook

from gdd_loader.extract.control_reader import (
    hash_archivo,
    leer_control,
    leer_estructura,
    leer_modificado_por,
)


def _crear(path, con_control=True, autor="Persona de Prueba"):
    wb = Workbook()
    ws = wb.active
    ws.title = "DetalleAtributos"
    ws.append(["codigo_dominio", " codigo_atributo ", None, "atributo", None, None])
    ws.append(["TST", 1, None, "x"])
    aux = wb.create_sheet("Lista De Referencia")
    aux.append(["id", "vals"])
    if con_control:
        c = wb.create_sheet("_Control")
        c.append(["clave", "valor"])
        c.append(["id_plantilla", "GDD-DOMINIO"])
        c.append(["version_plantilla", "1.0.0"])
        c.append([None, "ignorado"])
        c.append(["id_carga_base", 15])
        c.sheet_state = "veryHidden"
    wb.properties.lastModifiedBy = autor
    wb.save(path)
    return path


def test_leer_control_devuelve_claves_y_valores(tmp_path):
    archivo = _crear(tmp_path / "a.xlsx")
    assert leer_control(archivo) == {
        "id_plantilla": "GDD-DOMINIO", "version_plantilla": "1.0.0", "id_carga_base": "15",
    }


def test_leer_control_sin_hoja_devuelve_none(tmp_path):
    assert leer_control(_crear(tmp_path / "a.xlsx", con_control=False)) is None


def test_leer_estructura_excluye_control_y_limpia_encabezados(tmp_path):
    estructura = leer_estructura(_crear(tmp_path / "a.xlsx"))
    assert estructura == {
        "DetalleAtributos": ["codigo_dominio", "codigo_atributo", "", "atributo"],
        "Lista De Referencia": ["id", "vals"],
    }


def test_leer_modificado_por(tmp_path):
    assert leer_modificado_por(_crear(tmp_path / "a.xlsx")) == "Persona de Prueba"


def test_hash_archivo_es_sha256_del_contenido(tmp_path):
    archivo = tmp_path / "x.bin"
    archivo.write_bytes(b"contenido")
    assert hash_archivo(archivo) == hashlib.sha256(b"contenido").hexdigest()
