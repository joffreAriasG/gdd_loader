"""
Capa de extraccion: metadatos del archivo Excel que no son datos de negocio.

- hash del archivo completo (idempotencia y trazabilidad)
- hoja oculta `_Control` (identidad y version declaradas de la plantilla)
- encabezados de cada hoja (para calcular la huella de estructura)
- propiedad "Modificado por" del documento (dato INFORMATIVO de quien guardo
  el archivo por ultima vez; no sirve como control de acceso)

Usa openpyxl en modo solo lectura: no carga los datos de las hojas, solo lo
necesario.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

from openpyxl import load_workbook

from gdd_loader.domain.plantilla import HOJA_CONTROL, limpiar_encabezados


def hash_archivo(archivo: Path) -> str:
    h = hashlib.sha256()
    with open(archivo, "rb") as f:
        for bloque in iter(lambda: f.read(1024 * 1024), b""):
            h.update(bloque)
    return h.hexdigest()


def leer_control(archivo: Path) -> dict[str, str | None] | None:
    """Lee la hoja `_Control` (columna A = clave, columna B = valor, fila 1 =
    encabezado). Retorna None si la hoja no existe (plantillas anteriores a
    la Fase 1)."""
    wb = load_workbook(archivo, read_only=True, data_only=True)
    try:
        if HOJA_CONTROL not in wb.sheetnames:
            return None
        valores: dict[str, str | None] = {}
        for fila in wb[HOJA_CONTROL].iter_rows(min_row=2, max_col=2, values_only=True):
            clave, valor = (list(fila) + [None, None])[:2]
            if clave is None or str(clave).strip() == "":
                continue
            valores[str(clave).strip()] = None if valor is None else str(valor).strip()
        return valores
    finally:
        wb.close()


def leer_estructura(archivo: Path) -> dict[str, list[str]]:
    """Encabezados (fila 1) de todas las hojas del archivo, excepto `_Control`."""
    wb = load_workbook(archivo, read_only=True, data_only=True)
    try:
        estructura: dict[str, list[str]] = {}
        for ws in wb.worksheets:
            if ws.title == HOJA_CONTROL:
                continue
            primera = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            estructura[ws.title] = limpiar_encabezados(list(primera))
        return estructura
    finally:
        wb.close()


def leer_modificado_por(archivo: Path) -> str | None:
    wb = load_workbook(archivo, read_only=True)
    try:
        valor = wb.properties.lastModifiedBy or wb.properties.creator
        return str(valor).strip() or None if valor else None
    finally:
        wb.close()
