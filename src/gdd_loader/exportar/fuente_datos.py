"""
De donde salen los datos de una plantilla precargada (Fase 2).

La plantilla generada debe reflejar EXACTAMENTE la ultima carga exitosa
(MERGE_OK) del dominio. Dos fuentes, en orden de preferencia:

1. ARCHIVO: el .xlsx archivado de esa carga (carpeta Archivo). Exacto,
   incluido el orden de las filas. La purga nunca borra el archivo de la
   ultima MERGE_OK de un dominio.
2. STAGING: las filas de staging.* marcadas con ese id_carga. Tambien exacto
   en valores (staging es espejo del Excel), pero sin orden original: se
   ordena por la clave de la hoja. Solo sirve si NINGUNA carga posterior
   (fallida) reemplazo staging de ese dominio -- se verifica fila por fila.

No se reconstruye desde las tablas gdd.*: esa transformacion no es
reversible sin perdida (catalogos resueltos por codigo, agrupacion de las
columnas _foc, etc.).

Ambas fuentes devuelven lo mismo: hoja -> lista de filas como dict
{encabezado: valor}. `alinear_a_molde` las ordena segun las columnas del
molde vigente (por nombre, no por posicion), que es lo que permite pasar
datos de una version de plantilla a otra.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from gdd_loader.domain.plantilla import normalizar_encabezado
from gdd_loader.domain.sheet_config import SheetConfig

FilasPorHoja = dict[str, list[dict]]


class FuenteNoDisponibleError(Exception):
    pass


def _vacio(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def datos_desde_archivo(archivo: Path, hojas: list[str]) -> FilasPorHoja:
    """Filas no vacias de cada hoja, en el orden del archivo."""
    if not archivo.exists():
        raise FuenteNoDisponibleError(f"No existe el archivo archivado {archivo}")
    wb = load_workbook(archivo, read_only=True, data_only=True)
    try:
        datos: FilasPorHoja = {}
        for hoja in hojas:
            if hoja not in wb.sheetnames:
                raise FuenteNoDisponibleError(f"El archivo {archivo.name} no tiene la hoja '{hoja}'")
            filas = wb[hoja].iter_rows(values_only=True)
            encabezado = [normalizar_encabezado(v) for v in next(filas, ())]
            salida = []
            for fila in filas:
                if all(_vacio(v) for v in fila):
                    continue
                salida.append({h: v for h, v in zip(encabezado, fila) if h})
            datos[hoja] = salida
        return datos
    finally:
        wb.close()


def _clave_orden(valor) -> tuple:
    """Orden natural: 'CAC-10' despues de 'CAC-9'."""
    texto = "" if valor is None else str(valor)
    return tuple(int(p) if p.isdigit() else p for p in re.split(r"(\d+)", texto))


_ORDEN_HOJA = {
    "DetalleAtributos": "codigo_atributo",
    "MetadataTecnica": "codigo_dominio_atributo",
    "PlanDeRemediacion": "codigo_dominio_atributo",
    "Investigacion": "nombre",
    "Estructura": "codigo_plaza",
    "Respaldos": "tipo_respaldo",
}


def datos_desde_staging(staging_repo, hojas: list[SheetConfig], codigo_dominio: str,
                        id_carga: int) -> FilasPorHoja:
    datos: FilasPorHoja = {}
    for cfg in hojas:
        if cfg.columna_clave_prefijo_dominio:
            filas = staging_repo.leer_por_prefijo(cfg.tabla_staging, cfg.columna_clave, codigo_dominio)
        else:
            filas = staging_repo.leer(cfg.tabla_staging, "codigo_dominio", codigo_dominio)
        otras = {f.get("id_carga") for f in filas} - {id_carga}
        if otras:
            raise FuenteNoDisponibleError(
                f"{cfg.tabla_staging} tiene filas de otra carga ({sorted(map(str, otras))}); "
                f"staging ya no refleja la carga {id_carga}"
            )
        campo = _ORDEN_HOJA.get(cfg.nombre_hoja)
        if campo:
            filas = sorted(filas, key=lambda f: _clave_orden(f.get(campo)))
        datos[cfg.nombre_hoja] = [{k: v for k, v in f.items() if k in cfg.columnas} for f in filas]
    return datos


def alinear_a_molde(datos: FilasPorHoja, encabezados_molde: dict[str, list[str]]) -> tuple[dict[str, list[list]], list[str]]:
    """Convierte cada fila (dict) en una lista en el orden de columnas del
    molde. Una columna del molde sin dato en la fuente queda vacia (caso
    tipico: la version nueva agrego una columna). Una columna de la fuente
    que el molde ya no tiene se informa como advertencia (sus datos no pasan
    a la plantilla nueva)."""
    salida: dict[str, list[list]] = {}
    avisos: list[str] = []
    for hoja, columnas in encabezados_molde.items():
        filas = datos.get(hoja, [])
        presentes = set().union(*(f.keys() for f in filas)) if filas else set()
        nuevas = [c for c in columnas if c and c not in presentes]
        perdidas = sorted(presentes - set(columnas))
        if filas and nuevas:
            avisos.append(f"{hoja}: columnas nuevas sin datos de origen {nuevas} (quedan vacias)")
        if perdidas:
            avisos.append(f"{hoja}: columnas del origen que el molde ya no tiene {perdidas} (no se copian)")
        salida[hoja] = [[f.get(c) for c in columnas] for f in filas]
    return salida, avisos
