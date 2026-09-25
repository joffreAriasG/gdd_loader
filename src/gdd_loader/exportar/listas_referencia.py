"""
Listas de referencia de la plantilla (Fase 2): valores permitidos que los
Office Scripts usan para validar y que el loader luego resuelve contra los
catalogos de la BD.

Hojas del molde:
  - ListaDeReferencia (id, vals): un grupo por `id` (Criticidad, TipoAtributo...)
  - FuentesConsumo (coleccion, name): destinos de consumo
  - FuentesPrimarias (servidor, bdd): SIEMPRE se toma del molde (decision
    2026-09-25: la plantilla registrada contiene las fuentes reales de
    produccion; la BD puede tener servidores/bdd creados automaticamente por
    el flujo de fuente oficial, que no deben ofrecerse como opcion).

Modos (GDD_LISTAS_REFERENCIA):
  - MOLDE: las listas se copian del molde (ambiente de pruebas).
  - BD: los grupos con catalogo se toman de la BD; si un catalogo devuelve
    0 filas se conserva la lista del molde (nunca una lista vacia).

Correspondencia grupo -> catalogo:
  - CATALOGOS_SIMPLES: CONFIRMADOS, son la misma tabla/columna con la que el
    loader valida hoy (pipeline/carga_gdd.py, resolver_catalogo_controlado).
  - CATALOGOS_DATO_PERSONAL: HIPOTESIS de formato ("codigo. descripcion");
    `listas_cli comparar` la confirma o descarta contra los datos reales.
  - Grupos sin catalogo (DatoPersonal, AceptaNulos, TipoDatoCampo): del molde.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

HOJA_LISTAS = "ListaDeReferencia"
HOJA_CONSUMO = "FuentesConsumo"
HOJA_ERRORES = "Reporte_Errores"

CATALOGOS_SIMPLES: dict[str, tuple[str, str]] = {
    "Criticidad": ("cat_criticidad", "nombre_criticidad"),
    "TipoAtributo": ("cat_tipo_atributo", "nombre_tipo_atributo"),
    "CompartidoSimilar": ("cat_compartido_similar", "nombre_compartido_similar"),
    "TipoRespaldo": ("cat_tipo_respaldo", "nombre_tipo_respaldo"),
    "Dimension": ("cat_tipo_dimension", "nombre_tipo_dimension"),
    "TipoPlan": ("cat_tipo_plan", "nombre_tipo_plan"),
    "CategoriaPlan": ("cat_categoria_plan", "nombre_categoria_plan"),
    "SubcategoriaPlan": ("cat_sub_categoria_plan", "nombre_sub_categoria_plan"),
    "EstadoActual": ("cat_estado_plan", "nombre_estado_plan"),
}

# grupo -> (columna codigo, columna descripcion) en gdd.cat_dato_personal
CATALOGOS_DATO_PERSONAL: dict[str, tuple[str, str]] = {
    "CategoriaNivelUno": ("codigo_categoria_nivel_uno", "categoria_nivel_uno"),
    "TipoDatoPersonal": ("codigo_tipo_dato", "tipo_dato"),
    "Sensibilidad": ("codigo_sensibilidad", "sensibilidad"),
}


@dataclass
class ListasMolde:
    grupos: dict[str, list[str]]  # ListaDeReferencia, en el orden del molde
    consumo: list[tuple[str, str]]  # FuentesConsumo


@dataclass
class Diferencia:
    grupo: str
    origen_bd: str  # descripcion del catalogo o "sin catalogo"
    faltan_en_bd: list[str]
    solo_en_bd: list[str]

    @property
    def coincide(self) -> bool:
        return not self.faltan_en_bd and not self.solo_en_bd


def _txt(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def leer_listas_molde(molde: Path) -> ListasMolde:
    wb = load_workbook(molde, read_only=True, data_only=True)
    try:
        grupos: dict[str, list[str]] = {}
        if HOJA_LISTAS in wb.sheetnames:
            for fila in wb[HOJA_LISTAS].iter_rows(min_row=2, max_col=2, values_only=True):
                grupo, valor = (_txt(v) for v in (list(fila) + [None, None])[:2])
                if grupo and valor and valor not in grupos.setdefault(grupo, []):
                    grupos[grupo].append(valor)
        consumo: list[tuple[str, str]] = []
        if HOJA_CONSUMO in wb.sheetnames:
            for fila in wb[HOJA_CONSUMO].iter_rows(min_row=2, max_col=2, values_only=True):
                col, nombre = (_txt(v) for v in (list(fila) + [None, None])[:2])
                if col and nombre and (col, nombre) not in consumo:
                    consumo.append((col, nombre))
        return ListasMolde(grupos, consumo)
    finally:
        wb.close()


class ListasRepository:
    """Solo lectura sobre los catalogos gdd.*. Los nombres de tabla/columna
    vienen de las constantes de este modulo (lista blanca), nunca del usuario."""

    def __init__(self, engine):
        self._engine = engine

    def valores_simples(self, tabla: str, columna: str) -> list[str]:
        with self._engine.connect() as conn:
            filas = conn.execute(text(
                f"SELECT DISTINCT {columna} FROM gdd.{tabla} WHERE {columna} IS NOT NULL ORDER BY {columna}"
            )).all()
        return [str(f[0]).strip() for f in filas if _txt(f[0])]

    def valores_dato_personal(self, col_codigo: str, col_desc: str) -> list[str]:
        with self._engine.connect() as conn:
            filas = conn.execute(text(
                f"SELECT DISTINCT {col_codigo}, {col_desc} FROM gdd.cat_dato_personal "
                f"WHERE {col_codigo} IS NOT NULL"
            )).all()
        valores = {f"{str(c).strip()}. {str(d).strip()}" for c, d in filas if _txt(c) and _txt(d)}
        return sorted(valores, key=_orden_codigo)

    def fuentes_consumo(self) -> list[tuple[str, str]]:
        with self._engine.connect() as conn:
            filas = conn.execute(text(
                "SELECT DISTINCT b.nombre_bdd, t.nombre_tabla "
                "FROM gdd.base_datos_fuente_tabla t "
                "JOIN gdd.base_datos_fuente b ON b.id = t.id_bdd "
                "WHERE t.activo = 1 ORDER BY b.nombre_bdd, t.nombre_tabla"
            )).all()
        return [(str(b).strip(), str(t).strip()) for b, t in filas if _txt(b) and _txt(t)]


def _orden_codigo(valor: str) -> tuple:
    """'2.10. x' despues de '2.9. x'."""
    codigo = valor.split(" ", 1)[0].rstrip(".")
    return tuple(int(p) if p.isdigit() else 0 for p in codigo.split(".")), valor


def listas_bd(repo: ListasRepository) -> tuple[dict[str, list[str]], dict[str, str]]:
    """Valores de la BD por grupo, y descripcion del origen de cada grupo."""
    valores: dict[str, list[str]] = {}
    origen: dict[str, str] = {}
    for grupo, (tabla, col) in CATALOGOS_SIMPLES.items():
        valores[grupo] = repo.valores_simples(tabla, col)
        origen[grupo] = f"gdd.{tabla}.{col}"
    for grupo, (cod, desc) in CATALOGOS_DATO_PERSONAL.items():
        valores[grupo] = repo.valores_dato_personal(cod, desc)
        origen[grupo] = f"gdd.cat_dato_personal ({cod} + '. ' + {desc}) [formato por confirmar]"
    return valores, origen


def comparar(molde: ListasMolde, repo: ListasRepository) -> list[Diferencia]:
    valores, origen = listas_bd(repo)
    difs = []
    for grupo, vals_molde in molde.grupos.items():
        if grupo not in valores:
            difs.append(Diferencia(grupo, "sin catalogo (se usa el molde)", [], []))
            continue
        bd = valores[grupo]
        difs.append(Diferencia(grupo, origen[grupo],
                               [v for v in vals_molde if v not in bd],
                               [v for v in bd if v not in vals_molde]))
    bd_consumo = repo.fuentes_consumo()
    fmt = lambda t: f"{t[0]} / {t[1]}"  # noqa: E731
    difs.append(Diferencia(HOJA_CONSUMO, "gdd.base_datos_fuente + base_datos_fuente_tabla",
                           [fmt(t) for t in molde.consumo if t not in bd_consumo],
                           [fmt(t) for t in bd_consumo if t not in molde.consumo]))
    return difs


def construir_listas(molde: ListasMolde, repo: ListasRepository) -> tuple[dict[str, list[list]], list[str]]:
    """Modo BD: filas para ListaDeReferencia y FuentesConsumo. Se conservan
    los grupos del molde y su orden (los Office Scripts buscan por `id`); un
    grupo con catalogo vacio o sin catalogo se deja como en el molde."""
    valores, _ = listas_bd(repo)
    avisos = []
    filas_listas: list[list] = []
    for grupo, vals_molde in molde.grupos.items():
        vals = valores.get(grupo)
        if vals is None:
            vals = vals_molde
        elif not vals:
            avisos.append(f"{grupo}: el catalogo de la BD esta vacio; se conserva la lista del molde")
            vals = vals_molde
        filas_listas.extend([grupo, v] for v in vals)
    consumo = repo.fuentes_consumo()
    if not consumo:
        avisos.append(f"{HOJA_CONSUMO}: sin filas en la BD; se conserva la lista del molde")
        consumo = molde.consumo
    return {HOJA_LISTAS: filas_listas, HOJA_CONSUMO: [list(t) for t in consumo]}, avisos


def _sql_literal(valor: str) -> str:
    return "N'" + valor.replace("'", "''") + "'"


def script_semilla(molde: ListasMolde, difs: list[Diferencia]) -> str:
    """Script T-SQL (para el ambiente de PRUEBAS) que inserta en los
    catalogos simples los valores de la plantilla que faltan en la BD. No se
    ejecuta solo: se revisa y se corre a mano. Soporta id IDENTITY o no."""
    faltantes = {d.grupo: d.faltan_en_bd for d in difs if d.grupo in CATALOGOS_SIMPLES and d.faltan_en_bd}
    lineas = [
        "/* Semilla de catalogos para AMBIENTE DE PRUEBAS, generada desde la plantilla registrada.",
        "   Inserta solo valores que faltan (NOT EXISTS). Revisar antes de correr.",
        "   NO correr en produccion. Todo o nada: si algo falla se revierte completo. */",
        "SET XACT_ABORT ON;", "SET NOCOUNT ON;", "BEGIN TRANSACTION;", "",
    ]
    if not faltantes:
        lineas.append("-- No faltan valores en los catalogos simples.")
    for grupo, valores in faltantes.items():
        tabla, col = CATALOGOS_SIMPLES[grupo]
        lineas.append(f"-- {grupo} -> gdd.{tabla}.{col} ({len(valores)} valor(es))")
        for v in valores:
            lit = _sql_literal(v)
            lineas.append(
                f"IF NOT EXISTS (SELECT 1 FROM gdd.{tabla} WHERE {col} = {lit})\n"
                f"    IF COLUMNPROPERTY(OBJECT_ID('gdd.{tabla}'), 'id', 'IsIdentity') = 1\n"
                f"        INSERT INTO gdd.{tabla} ({col}) VALUES ({lit});\n"
                f"    ELSE\n"
                f"        INSERT INTO gdd.{tabla} (id, {col}) SELECT ISNULL(MAX(id), 0) + 1, {lit} FROM gdd.{tabla};"
            )
        lineas.append("")
    otros = [d for d in difs if d.grupo not in CATALOGOS_SIMPLES and d.faltan_en_bd]
    if otros:
        lineas.append("/* No incluidos (requieren datos que la plantilla no trae; cargar a mano si hace falta):")
        for d in otros:
            lineas.append(f"   - {d.grupo}: {len(d.faltan_en_bd)} valor(es) faltantes ({d.origen_bd})")
        lineas.append("*/")
    lineas += ["", "COMMIT TRANSACTION;"]
    return "\n".join(lineas) + "\n"
